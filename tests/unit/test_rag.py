import pytest
from unittest.mock import MagicMock, patch
from sentence_transformers import SentenceTransformer

from io import BytesIO
from docx import Document as DocxDocument
import openpyxl

from app.rag.reranker import Reranker
from app.rag.ingestion import DocumentIngester
from app.config import settings


def test_rerank_empty_candidates():
    """Test that rerank returns empty list when no candidates are provided."""
    mock_model = MagicMock()
    mock_model.predict.return_value = [0.9, 0.1, 0.5]
    reranker = Reranker(model=mock_model)
    res = reranker.rerank(query='', candidates=[], top_k=settings.TOP_K)
    assert res == []


def test_rerank_returns_top_k():
    """Test that rerank returns exactly top_k results when candidates exceed top_k."""
    mock_model = MagicMock()
    mock_model.predict.return_value = [0.9, 0.7, 0.5]
    candidates = [
        ("Текст чанка 1", 0.9),
        ("Текст чанка 2", 0.7),
        ("Текст чанка 3", 0.5)
    ]
    reranker = Reranker(model=mock_model)
    res = reranker.rerank(query='', candidates=candidates, top_k=2)
    assert len(res) == 2


def test_rerank_sorted_by_score():
    """Test that rerank returns candidates sorted by descending cross-encoder score."""
    mock_model = MagicMock()
    mock_model.predict.return_value = [0.5, 0.3, 0.9]
    candidates = [("text_a", 0.5), ("text_b", 0.3), ("text_c", 0.9)]
    reranker = Reranker(mock_model)
    res = reranker.rerank(query='', candidates=candidates, top_k=3)
    assert res[0] == "text_c"


def test_load_txt():
    """Test that _load_txt correctly decodes bytes content to a list with one string."""
    with patch('app.rag.ingestion.AsyncQdrantClient'):
        ingester = DocumentIngester(
            content=b"hello world",
            filename="test.txt",
            user_id="test_user",
            embedding_model=MagicMock(spec=SentenceTransformer)
        )
        result = ingester._load_txt()
        assert result == ["hello world"]


def test_load_docx():
    """Test that _load_docx extracts non-empty paragraphs from a DOCX file in memory."""
    paragraphs = ['paragraph_1', 'paragraph_2', 'paragraph_3']
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = BytesIO()
    doc.save(buf)
    sample = buf.getvalue()
    with patch('app.rag.ingestion.AsyncQdrantClient'):  # ← обязательно
        ingester = DocumentIngester(
            content=sample,
            filename="test.docx",
            user_id="test_user",
            embedding_model=MagicMock(spec=SentenceTransformer)
        )
        result = ingester._load_docx()
        assert result == ["paragraph_1", "paragraph_2", "paragraph_3"]


def test_load_xlsx():
    """Test that _load_xlsx converts Excel rows to text representations."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Item', 'Sales'])
    ws.append(['Phone case', 100])
    buf = BytesIO()
    wb.save(buf)
    sample = buf.getvalue()
    with patch('app.rag.ingestion.AsyncQdrantClient'):  # ← обязательно
        ingester = DocumentIngester(
            content=sample,
            filename="test.xlsx",
            user_id="test_user",
            embedding_model=MagicMock(spec=SentenceTransformer)
        )
        result = ingester._load_xlsx()
        assert len(result) == 1
        assert "Phone case" in result[0]
        assert "100" in result[0]


@pytest.mark.asyncio
async def test_ingest_unsupported_type():
    """Test that ingest raises ValueError for unsupported file extensions."""
    with patch('app.rag.ingestion.AsyncQdrantClient'):
        ingester = DocumentIngester(
            content=b"data",
            filename="file.pdf",
            user_id="test_user",
            embedding_model=MagicMock(spec=SentenceTransformer)
        )
        with pytest.raises(ValueError):
            await ingester.ingest()